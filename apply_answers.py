"""
Apply answers from intermediate JSON or YAML files to Assessment Workbook (AWB) docx files.
Supports reproducible, idempotent marking so fields can be re-identified and updated repeatedly.
"""
import os
import sys
import json
import yaml
import copy
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
W14_NS = 'http://schemas.microsoft.com/office/word/2010/wordml'

def get_field_by_mark(doc, field_id, field_type):
    """
    Search document for field by its embedded unique mark:
    - Text fields: w:name in w:ffData
    - Checkboxes: w:tag in w:sdtPr
    """
    if field_type == "text":
        for ff in doc._element.iter(f'{{{W_NS}}}ffData'):
            name_el = ff.find(f'{{{W_NS}}}name')
            if name_el is not None and name_el.get(f'{{{W_NS}}}val') == field_id:
                # Find parent paragraph
                curr = ff
                while curr is not None and curr.tag != f'{{{W_NS}}}p':
                    curr = curr.getparent()
                return curr
    elif field_type == "checkbox":
        for sdt in doc._element.iter(f'{{{W_NS}}}sdt'):
            tag_el = sdt.find(f'.//{{{W_NS}}}tag')
            if tag_el is not None and tag_el.get(f'{{{W_NS}}}val') == field_id:
                return sdt
    return None

def set_sdt_tag(sdt_elem, field_id):
    """Ensure w:tag with field_id exists inside w:sdtPr."""
    sdtPr = sdt_elem.find(f'{{{W_NS}}}sdtPr')
    if sdtPr is None:
        sdtPr = OxmlElement('w:sdtPr')
        sdt_elem.insert(0, sdtPr)
    tag_el = sdtPr.find(f'{{{W_NS}}}tag')
    if tag_el is None:
        tag_el = OxmlElement('w:tag')
        sdtPr.append(tag_el)
    tag_el.set(f'{{{W_NS}}}val', field_id)

def set_ffdata_name(p_elem, field_id, ff_index=0):
    """Ensure w:name with field_id is set in w:ffData."""
    ff_list = p_elem.findall(f'.//{{{W_NS}}}ffData')
    if ff_index < len(ff_list):
        ff = ff_list[ff_index]
        name_el = ff.find(f'{{{W_NS}}}name')
        if name_el is None:
            name_el = OxmlElement('w:name')
            ff.insert(0, name_el)
        name_el.set(f'{{{W_NS}}}val', field_id)

def update_text_field(p_elem, text_value, field_id, ff_index=0):
    """
    Update text form field in paragraph.
    Replaces runs between separate and end with new text runs.
    """
    # Set permanent mark
    set_ffdata_name(p_elem, field_id, ff_index)
    
    # Locate all begin, separate, end markers
    w_r = f'{{{W_NS}}}r'
    w_fldChar = f'{{{W_NS}}}fldChar'
    w_type = f'{{{W_NS}}}fldCharType'
    
    runs = [c for c in p_elem if c.tag == w_r]
    
    # Find the separate and end markers corresponding to ff_index
    current_begin = -1
    sep_idx = -1
    end_idx = -1
    target_begin_count = 0
    
    for i, r in enumerate(runs):
        for child in r:
            if child.tag == w_fldChar:
                f_type = child.get(w_type)
                if f_type == "begin":
                    if target_begin_count == ff_index:
                        current_begin = i
                    target_begin_count += 1
                elif f_type == "separate" and current_begin != -1 and sep_idx == -1:
                    sep_idx = i
                elif f_type == "end" and sep_idx != -1:
                    end_idx = i
                    break
        if end_idx != -1:
            break
            
    if sep_idx == -1 or end_idx == -1 or end_idx <= sep_idx:
        return False

    # Get sample rPr for styling
    rPr = None
    if sep_idx > 0:
        rPr_elem = runs[sep_idx].find(f'{{{W_NS}}}rPr')
        if rPr_elem is not None:
            rPr = copy.deepcopy(rPr_elem)
    if rPr is None:
        # Minor minor default run properties
        rPr = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(f'{{{W_NS}}}val', '404040')
        rPr.append(color)

    # Remove existing runs between sep_idx and end_idx
    runs_to_remove = runs[sep_idx + 1 : end_idx]
    for r in runs_to_remove:
        p_elem.remove(r)
        
    # Create new run(s) with text_value
    text_str = str(text_value) if text_value is not None else ""
    # If text has multiple paragraphs or newlines, handle clean insertion
    lines = text_str.split('\n')
    
    # Find insertion point: directly after sep_run
    sep_run = runs[sep_idx]
    insert_pos = p_elem.index(sep_run) + 1
    
    new_r = OxmlElement('w:r')
    if rPr is not None:
        new_r.append(copy.deepcopy(rPr))
        
    for line_idx, line in enumerate(lines):
        t_el = OxmlElement('w:t')
        t_el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        t_el.text = line
        new_r.append(t_el)
        if line_idx < len(lines) - 1:
            br_el = OxmlElement('w:br')
            new_r.append(br_el)
            
    p_elem.insert(insert_pos, new_r)
    return True

def update_checkbox(sdt_elem, is_checked, field_id):
    """
    Update checkbox SDT:
    - Sets w:tag to field_id
    - Sets w14:checked to 1 or 0
    - Updates sdtContent run text to ☒ or ☐
    """
    # 1. Permanent mark
    set_sdt_tag(sdt_elem, field_id)
    
    # 2. Update checked attribute
    cb = sdt_elem.find(f'.//{{{W14_NS}}}checkbox')
    if cb is not None:
        checked_el = cb.find(f'{{{W14_NS}}}checked')
        if checked_el is None:
            checked_el = OxmlElement('w14:checked')
            cb.append(checked_el)
        checked_el.set(f'{{{W14_NS}}}val', '1' if is_checked else '0')
        
    # 3. Update text in sdtContent
    sdtContent = sdt_elem.find(f'{{{W_NS}}}sdtContent')
    if sdtContent is not None:
        t_el = sdtContent.find(f'.//{{{W_NS}}}t')
        if t_el is not None:
            t_el.text = '☒' if is_checked else '☐'
        else:
            r = OxmlElement('w:r')
            t = OxmlElement('w:t')
            t.text = '☒' if is_checked else '☐'
            r.append(t)
            sdtContent.append(r)
    return True

def apply_answers(answers_file, input_docx, output_docx):
    print(f"Applying answers from {answers_file} to {input_docx}...")
    
    # Load answers
    if answers_file.endswith('.json'):
        with open(answers_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        with open(answers_file, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)
            
    doc = docx.Document(input_docx)
    fields = data.get('fields', [])
    
    applied_count = 0
    skipped_count = 0
    not_found_count = 0
    
    for f in fields:
        fid = f['field_id']
        ftype = f['type']
        fval = f.get('value')
        t_idx = f.get('table_idx')
        r_idx = f.get('row_idx')
        c_idx = f.get('col_idx')
        idx_in_cell = f.get('index_in_cell', 0)
        
        # 1. Try finding by permanent mark first
        target_elem = get_field_by_mark(doc, fid, ftype)
        
        # 2. If not found by mark, locate by table coordinates
        p_ff_index = 0
        if target_elem is None:
            if t_idx is not None and t_idx < len(doc.tables):
                t = doc.tables[t_idx]
                if r_idx is not None and r_idx < len(t.rows):
                    r = t.rows[r_idx]
                    if c_idx is not None and c_idx < len(r.cells):
                        cell = r.cells[c_idx]
                        if ftype == "text":
                            # Gather all form fields across all paragraphs in the cell in document order
                            all_ff = []
                            for p_elem in cell._element.iter(f'{{{W_NS}}}p'):
                                ff_in_p = p_elem.findall(f'.//{{{W_NS}}}ffData')
                                for p_ff_i in range(len(ff_in_p)):
                                    all_ff.append((p_elem, p_ff_i))
                            if idx_in_cell < len(all_ff):
                                target_elem, p_ff_index = all_ff[idx_in_cell]
                        elif ftype == "checkbox":
                            sdts = cell._element.findall(f'.//{{{W_NS}}}sdt')
                            if idx_in_cell < len(sdts):
                                target_elem = sdts[idx_in_cell]
                                
        if target_elem is None:
            not_found_count += 1
            continue
            
        # 3. Apply the value
        if ftype == "text":
            if fval is not None and str(fval).strip():
                update_text_field(target_elem, fval, fid, p_ff_index)
                applied_count += 1
            else:
                set_ffdata_name(target_elem, fid, p_ff_index)
                skipped_count += 1
        elif ftype == "checkbox":
            is_checked = bool(fval)
            update_checkbox(target_elem, is_checked, fid)
            applied_count += 1
            
    print(f"Applied: {applied_count}, Blank/Skipped: {skipped_count}, Not Found: {not_found_count}")
    doc.save(output_docx)
    print(f"Successfully saved filled document to {output_docx}")

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Apply answers to AWB docx documents.")
    parser.add_argument("--answers", help="Path to YAML or JSON answer file")
    parser.add_argument("--input", help="Path to input AWB docx")
    parser.add_argument("--output", help="Path to output filled docx")
    parser.add_argument("--unit", help="Specific unit code to apply (e.g. CHCCCS040, CHCCOM005, HLTINF006, CHCAGE013)")
    args = parser.parse_args()
    
    BATCH_CONFIGS = [
        # Initial 2 units
        ("CHCDIS020", "answers_CHCDIS020.yaml", "Assignment Materials-20260914/CHCDIS020-AWB-F-v1.0.docx", "Assignment Materials-20260914/CHCDIS020-AWB-Filled.docx"),
        ("CHCPAL003", "answers_CHCPAL003.yaml", "Assignment Materials-20260915/CHCPAL003-AWB-F-v1.0 .docx", "Assignment Materials-20260915/CHCPAL003-AWB-Filled.docx"),
        # Group 1 units
        ("CHCCCS040", "answers_CHCCCS040.yaml", "Assignment Materials-20260915 (3)/CHCCCS040-AWB-F-v1.0.docx", "Assignment Materials-20260915 (3)/CHCCCS040-AWB-Filled.docx"),
        ("CHCCOM005", "answers_CHCCOM005.yaml", "Assignment Materials-20260915 (5)/CHCCOM005-AWB-F-v1.1.docx", "Assignment Materials-20260915 (5)/CHCCOM005-AWB-Filled.docx"),
        ("HLTINF006", "answers_HLTINF006.yaml", "Assignment Materials-20260915 (8)/HLTINF006-AWB-F-v1.0.docx", "Assignment Materials-20260915 (8)/HLTINF006-AWB-Filled.docx"),
        ("CHCAGE013", "answers_CHCAGE013.yaml", "Assignment Materials-20260915 (11)/CHCAGE013-AWB-F-v1.0.docx", "Assignment Materials-20260915 (11)/CHCAGE013-AWB-Filled.docx"),
    ]

    if args.answers and args.input and args.output:
        apply_answers(args.answers, args.input, args.output)
    elif args.unit:
        matched = [c for c in BATCH_CONFIGS if c[0].lower() == args.unit.lower()]
        if matched:
            _, ans, inp, out = matched[0]
            apply_answers(ans, inp, out)
        else:
            print(f"Unknown unit code: {args.unit}")
    else:
        for code, ans, inp, out in BATCH_CONFIGS:
            if os.path.exists(ans) and os.path.exists(inp):
                print(f"\n{'='*25} Applying {code} {'='*25}")
                apply_answers(ans, inp, out)

if __name__ == "__main__":
    main()
