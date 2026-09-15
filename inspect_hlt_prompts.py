import docx

doc = docx.Document("8. HLTINF006 - Apply basic principles and practices of infection prevention and control/HLTINF006-AWB-F-v1.0.docx")
tables = [28, 30, 36, 37, 38, 39, 40, 41, 42, 43, 49, 50, 51, 58, 59, 78]

for t_i in tables:
    t = doc.tables[t_i]
    p0 = " ".join([c.text.strip().replace('\u2002', '') for c in t.rows[0].cells if c.text.strip()])
    p_prev = ""
    if t_i > 0:
        prev_t = doc.tables[t_i-1]
        p_prev = " ".join([c.text.strip().replace('\u2002', '') for c in prev_t.rows[0].cells if c.text.strip()])[:60]
    print(f"Table {t_i:2d} ({len(t.rows)} rows): R0='{p0[:60]}' | Prev='{p_prev}'")
