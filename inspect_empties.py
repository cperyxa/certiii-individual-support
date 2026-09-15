import docx
import yaml

def inspect_unit_empties(yaml_path, awb_path, ag_path):
    print(f"\n{'='*30} INSPECTING {yaml_path} {'='*30}")
    with open(yaml_path, 'r', encoding='utf-8') as f:
        data = yaml.safe_load(f)
        
    awb = docx.Document(awb_path)
    ag = docx.Document(ag_path)
    
    empty_cand = [f for f in data['fields'] if f['type'] == 'text' and f['section'] != 'Assessor Section' and (not f['value'] or not str(f['value']).strip())]
    print(f"Total empty: {len(empty_cand)}")
    
    seen_tables = set()
    for ec in empty_cand:
        t_idx = ec['table_idx']
        if t_idx not in seen_tables:
            seen_tables.add(t_idx)
            # Find prompt text in AWB
            t = awb.tables[t_idx]
            awb_txt = " ".join([c.text.strip().replace('\u2002', '') for r in t.rows for c in r.cells if c.text.strip()])
            print(f"\n--- AWB Table {t_idx} (rows: {len(t.rows)}) ---")
            print(f"Text snippet: {awb_txt[:120]}")
            
            # Search AG for matching question
            for i, ag_t in enumerate(ag.tables):
                ag_txt = " ".join([c.text for r in ag_t.rows for c in r.cells])
                # Check keyword overlap
                words = [w for w in awb_txt.lower().split() if len(w) > 4][:5]
                if any(w in ag_txt.lower() for w in words):
                    # Potential match
                    first_cell = ag_t.rows[0].cells[0].text.strip().replace('\n', ' ')[:80]
                    print(f"  -> Potential AG Table {i}: {first_cell}")
                    break

inspect_unit_empties("answers_CHCCCS040.yaml", "Assignment Materials-20260915 (3)/CHCCCS040-AWB-F-v1.0.docx", "Assignment Materials-20260915 (3)/CHCCCS040-AG-F-v1.0.docx")
inspect_unit_empties("answers_CHCCOM005.yaml", "Assignment Materials-20260915 (5)/CHCCOM005-AWB-F-v1.1.docx", "Assignment Materials-20260915 (5)/CHCCOM005-AG-F-v1.1.docx")
inspect_unit_empties("answers_HLTINF006.yaml", "Assignment Materials-20260915 (8)/HLTINF006-AWB-F-v1.0.docx", "Assignment Materials-20260915 (8)/HLTINF006-AG-F-v1.0.docx")
inspect_unit_empties("answers_CHCAGE013.yaml", "Assignment Materials-20260915 (11)/CHCAGE013-AWB-F-v1.0.docx", "Assignment Materials-20260915 (11)/CHCAGE013-AG-F-v1.1.docx")
