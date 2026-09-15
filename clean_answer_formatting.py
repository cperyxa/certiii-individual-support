"""
Module to format filled Assessment Workbooks:
1. Hide the grey background on all candidate answers (unlinks FORMTEXT fields into clean text runs and adds doNotShadeFormData to settings).
2. Hide border lines around open-ended question answer boxes (sets tcBorders to nil).
3. Keep table grid lines intact for structured multi-column tables.
4. Leave Assessor Sections completely untouched.
"""
import os
import sys
import docx
from docx.oxml import OxmlElement

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def unlink_fields_in_p(p_elem):
    """Convert form field into normal text run, removing grey field shading."""
    children = list(p_elem)
    in_field = False
    in_result = False
    to_remove = []
    
    for child in children:
        if child.tag == f'{{{W_NS}}}r':
            fldChars = child.findall(f'.//{{{W_NS}}}fldChar')
            for fc in fldChars:
                ftype = fc.get(f'{{{W_NS}}}fldCharType')
                if ftype == 'begin':
                    in_field = True
                    in_result = False
                    to_remove.append(child)
                elif ftype == 'separate':
                    in_result = True
                    to_remove.append(child)
                elif ftype == 'end':
                    in_field = False
                    in_result = False
                    to_remove.append(child)
            if in_field and not in_result:
                if child not in to_remove:
                    to_remove.append(child)
                    
    for rem in to_remove:
        p_elem.remove(rem)

def hide_cell_borders(cell):
    """Set top, left, bottom, right borders to nil."""
    tcPr = cell._tc.find(f'{{{W_NS}}}tcPr')
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        cell._tc.insert(0, tcPr)
    tcBorders = tcPr.find(f'{{{W_NS}}}tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for b_name in ['top', 'left', 'bottom', 'right']:
        b_el = tcBorders.find(f'{{{W_NS}}}{b_name}')
        if b_el is None:
            b_el = OxmlElement(f'w:{b_name}')
            tcBorders.append(b_el)
        b_el.set(f'{{{W_NS}}}val', 'nil')

def hide_grey_background_and_box_lines(doc):
    """
    Format document:
    - Hide grey field shading on candidate answers
    - Hide borders on open-ended question answer boxes
    - Keep structured table grids intact
    - Keep Assessor Sections intact
    """
    # 1. Ensure settings.xml has doNotShadeFormData
    try:
        settings_elem = doc.settings._element
        if settings_elem.find(f'{{{W_NS}}}doNotShadeFormData') is None:
            dn_shade = OxmlElement('w:doNotShadeFormData')
            settings_elem.append(dn_shade)
    except Exception:
        pass

    assessor_start = None
    n_tables = len(doc.tables)
    for i, t in enumerate(doc.tables):
        if i < n_tables // 2:
            continue
        txt = " ".join(c.text.strip().lower() for r in t.rows for c in r.cells)
        if (('to the assessor' in txt and 'completed assessing' in txt) or 
            ('record of assessment' in txt and any(k in txt for k in ['use only', 'candidate’s name', "candidate's name", 'outcome']))):
            assessor_start = i
            break

    open_boxes_hidden = 0
    fields_unshaded = 0

    for t_idx, t in enumerate(doc.tables):
        if assessor_start is not None and t_idx >= assessor_start:
            continue
            
        row_distinct = []
        for r in t.rows:
            seen = set(id(c._tc) for c in r.cells)
            row_distinct.append(len(seen))
        max_dist = max(row_distinct) if row_distinct else 0

        is_q_container = (max_dist == 1) or (max_dist == 2 and any(k in t.rows[0].cells[0].text.lower() for k in ['answer the following', 'read the scenario', 'which standard', 'scenario', 'task', 'question']))
        
        if max_dist == 1:
            tblPr = t._tbl.find(f'{{{W_NS}}}tblPr')
            if tblPr is not None:
                tblBorders = tblPr.find(f'{{{W_NS}}}tblBorders')
                if tblBorders is not None:
                    for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                        b_el = tblBorders.find(f'{{{W_NS}}}{b_name}')
                        if b_el is not None:
                            b_el.set(f'{{{W_NS}}}val', 'nil')

        for r_idx, r in enumerate(t.rows):
            seen_tc = set()
            is_single_merged_row = (row_distinct[r_idx] == 1)
            
            for c_idx, c in enumerate(r.cells):
                if id(c._tc) in seen_tc:
                    continue
                seen_tc.add(id(c._tc))
                
                ffs = c._tc.findall(f'.//{{{W_NS}}}ffData')
                if ffs:
                    fields_unshaded += len(ffs)
                    for p in c.paragraphs:
                        unlink_fields_in_p(p._element)
                        
                    if is_single_merged_row or (is_q_container and r_idx > 0):
                        hide_cell_borders(c)
                        open_boxes_hidden += 1
                        
    return fields_unshaded, open_boxes_hidden

def format_document(docx_path):
    doc = docx.Document(docx_path)
    f_count, b_count = hide_grey_background_and_box_lines(doc)
    doc.save(docx_path)
    print(f"Formatted {docx_path}: {f_count} fields unshaded, {b_count} open-ended box borders hidden.")

def format_all_documents():
    all_filled_files = [
        # Initial 2
        "Assignment Materials-20260914/CHCDIS020-AWB-Filled.docx",
        "Assignment Materials-20260915/CHCPAL003-AWB-Filled.docx",
        # Group 1 (4)
        "Assignment Materials-20260915 (3)/CHCCCS040-AWB-Filled.docx",
        "Assignment Materials-20260915 (5)/CHCCOM005-AWB-Filled.docx",
        "Assignment Materials-20260915 (8)/HLTINF006-AWB-Filled.docx",
        "Assignment Materials-20260915 (11)/CHCAGE013-AWB-Filled.docx",
        # Group 2 (5)
        "Assignment Materials-20260915 (1)/CHCCCS031-AWB-Part A-Filled.docx",
        "Assignment Materials-20260915 (1)/CHCCCS031-AWB-Part B-Filled.docx",
        "Assignment Materials-20260915 (2)/CHCCCS038-AWB-Filled.docx",
        "Assignment Materials-20260915 (4)/CHCCCS041-AWB-Filled.docx",
        "Assignment Materials-20260915 (6)/CHCDIV001-AWB-Filled.docx",
        # Group 3 (5)
        "Assignment Materials-20260915 (7)/CHCLEG001-AWB-Filled.docx",
        "Assignment Materials-20260915 (9)/HLTWHS002-AWB-Filled.docx",
        "Assignment Materials-20260915 (10)/CHCAGE011-AWB-Filled.docx",
        "Assignment Materials-20260915 (12)/CHCDIS011-AWB-Filled.docx",
        "Assignment Materials-20260915 (13)/CHCDIS012-AWB-Filled.docx",
    ]
    
    print(f"=== Formatting all {len(all_filled_files)} Assessment Workbooks ===")
    for path in all_filled_files:
        if os.path.exists(path):
            format_document(path)
        else:
            print(f"WARNING: {path} not found!")

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] != "--all":
        format_document(sys.argv[1])
    else:
        format_all_documents()
