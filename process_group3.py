"""
Process and generate baseline YAML answers for Group 3 units:
- CHCLEG001
- HLTWHS002
- CHCAGE011
- CHCDIS011
- CHCDIS012
"""
import os
import sys
import re
import yaml
import docx

from group3_overrides import CANDIDATE_DETAILS

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
W14_NS = 'http://schemas.microsoft.com/office/word/2010/wordml'

def clean_text(s):
    if not s: return ""
    return " ".join(str(s).replace('\u2002', ' ').split())

def process_file(unit_code, awb_path, output_yaml, specific_overrides=None):
    print(f"\n{'='*30} Processing {unit_code} {'='*30}")
    print(f"AWB: {awb_path}")
    
    awb_doc = docx.Document(awb_path)
    
    data = {
        "unit": unit_code,
        "candidate_details": CANDIDATE_DETAILS,
        "fields": []
    }
    
    cand_start = None
    assessor_start = None
    for i, t in enumerate(awb_doc.tables):
        txt = " ".join(c.text.strip() for r in t.rows for c in r.cells).lower()
        if 'to the candidate' in txt and cand_start is None:
            cand_start = i
        if ('to the assessor' in txt or 'record of assessment' in txt) and assessor_start is None and cand_start is not None:
            assessor_start = i

    for t_idx, t in enumerate(awb_doc.tables):
        t_text = " ".join([c.text.strip().replace('\u2002', '') for r in t.rows for c in r.cells if c.text.strip()])
        t_text_lower = t_text.lower()
        
        # Section detection
        if t_idx in [1, 2, 3] and ("workbook" in t_text_lower or "candidate declaration" in t_text_lower or "title" in t_text_lower):
            section = "Cover Sheet"
        elif "preliminary task" in t_text_lower or (cand_start is not None and t_idx < 5 and "state" in t_text_lower):
            section = "Preliminary Task"
        elif "to be completed by the manager" in t_text_lower or (assessor_start is not None and t_idx >= assessor_start) or "record of assessment" in t_text_lower:
            section = "Assessor Section"
        elif cand_start is not None and t_idx >= cand_start:
            section = "Candidate Checklist"
        elif any(k in t_text_lower for k in ["case study", "workplace assessment", "incident report", "client information", "observation"]):
            section = "Practical Assessment"
        else:
            section = "Knowledge Assessment"
            
        m_q = re.search(r'(Question\s+\d+|Task\s+[\d\.]+|Complete the table|Answer the following|Read the scenario|Briefly define|List down|Identify)', t_text, re.IGNORECASE)
        q_ref = m_q.group(1) if m_q else f"Table_{t_idx}"
        
        for r_idx, r in enumerate(t.rows):
            row_header = r.cells[0].text.strip().replace('\u2002', '').strip() if len(r.cells) > 0 else ""
            seen_tc = set()
            for c_idx, c in enumerate(r.cells):
                if id(c._tc) in seen_tc:
                    continue
                seen_tc.add(id(c._tc))
                
                col_header = ""
                if len(t.rows) > 0 and c_idx < len(t.rows[0].cells):
                    col_header = t.rows[0].cells[c_idx].text.strip().replace('\u2002', '').strip()
                cell_elem = c._element
                
                ff_list = cell_elem.findall(f'.//{{{W_NS}}}ffData')
                cb_list = cell_elem.findall(f'.//{{{W14_NS}}}checkbox')
                
                if not ff_list and not cb_list:
                    continue
                    
                # Handle checkboxes
                for cb_i, cb in enumerate(cb_list):
                    curr = cb
                    while curr is not None and not curr.tag.endswith('sdt'):
                        curr = curr.getparent()
                    sdt_id = ""
                    if curr is not None:
                        id_el = curr.find(f'.//{{{W_NS}}}id')
                        if id_el is not None:
                            sdt_id = id_el.get(f'{{{W_NS}}}val')
                            
                    cb_label = c.text.strip().replace('\n', ' ')
                    if not cb_label and c_idx + 1 < len(r.cells):
                        cb_label = r.cells[c_idx + 1].text.strip()
                    if not cb_label:
                        cb_label = f"Row {r_idx} Col {c_idx}"
                        
                    cb_value = False
                    if section == "Preliminary Task":
                        if "new south wales" in cb_label.lower():
                            cb_value = True
                    elif section == "Candidate Checklist":
                        if re.search(r'\bno\b', cb_label.lower()):
                            cb_value = False
                        else:
                            cb_value = True
                    elif section == "Practical Assessment":
                        if any(w in cb_label.lower() for w in ['bruise', 'burn', 'fear', 'withdrawn', 'injury', 'depression', 'anxiety', 'yes', 'tick', 'satisfactory', 'completed']):
                            cb_value = True
                    elif section == "Assessor Section":
                        cb_value = False
                            
                    field_id = f"{unit_code}_CB_{t_idx}_{r_idx}_{c_idx}_{cb_i+1}"
                    data["fields"].append({
                        "field_id": field_id,
                        "type": "checkbox",
                        "section": section,
                        "question_ref": q_ref,
                        "label": cb_label[:100],
                        "table_idx": t_idx,
                        "row_idx": r_idx,
                        "col_idx": c_idx,
                        "index_in_cell": cb_i,
                        "sdt_id": sdt_id,
                        "value": cb_value
                    })
                    
                # Handle text fields
                for ff_i, ff in enumerate(ff_list):
                    curr = ff
                    while curr is not None and not curr.tag.endswith('p'):
                        curr = curr.getparent()
                    para_id = curr.get(f'{{{W14_NS}}}paraId', '') if curr is not None else ""
                    
                    prompt_snippet = c.text.strip().replace('\u2002', '').strip()
                    if not prompt_snippet:
                        prompt_snippet = f"{row_header} - {col_header}".strip(" -")
                    if not prompt_snippet:
                        prompt_snippet = f"Table {t_idx} Row {r_idx} Col {c_idx}"
                        
                    answer = ""
                    if section == "Cover Sheet":
                        # Cover Sheet field mappings
                        p_low = prompt_snippet.lower()
                        if "first and last name" in p_low or r_idx == 2 and c_idx in [1, 2]:
                            answer = CANDIDATE_DETAILS["candidate_name"]
                        elif "phone" in p_low or r_idx == 3 and c_idx in [1, 2]:
                            answer = CANDIDATE_DETAILS["candidate_phone"]
                        elif "email" in p_low or r_idx == 4 and c_idx in [1, 2]:
                            answer = CANDIDATE_DETAILS["candidate_email"]
                        elif "date signed" in p_low or "date" in p_low:
                            answer = CANDIDATE_DETAILS["date"]
                        elif "signature" in p_low or "name:" in p_low:
                            answer = CANDIDATE_DETAILS["candidate_name"]
                        elif t_idx in [2, 3]:
                            if c_idx == 0 or "name" in p_low: answer = CANDIDATE_DETAILS["candidate_name"]
                            elif c_idx == 1 or "signature" in p_low: answer = CANDIDATE_DETAILS["candidate_name"]
                            elif c_idx == 2 or "date" in p_low: answer = CANDIDATE_DETAILS["date"]
                    elif section == "Candidate Checklist":
                        if "date" in prompt_snippet.lower():
                            answer = CANDIDATE_DETAILS["date"]
                        elif "signature" in prompt_snippet.lower() or "name" in prompt_snippet.lower():
                            answer = CANDIDATE_DETAILS["candidate_name"]
                    elif section == "Assessor Section":
                        if "candidate’s name" in prompt_snippet.lower() or "candidate's name" in prompt_snippet.lower():
                            answer = CANDIDATE_DETAILS["candidate_name"]
                        elif "rto name" in prompt_snippet.lower():
                            answer = CANDIDATE_DETAILS["rto"]
                        elif "rto contact" in prompt_snippet.lower():
                            answer = CANDIDATE_DETAILS["rto_phone"]
                        elif "rto email" in prompt_snippet.lower():
                            answer = CANDIDATE_DETAILS["rto_email"]
                    else:
                        # Check overrides
                        if specific_overrides and (t_idx, r_idx, c_idx, ff_i) in specific_overrides:
                            answer = specific_overrides[(t_idx, r_idx, c_idx, ff_i)]
                        elif specific_overrides and (t_idx, r_idx, c_idx) in specific_overrides:
                            answer = specific_overrides[(t_idx, r_idx, c_idx)]
                        elif specific_overrides and (t_idx, r_idx) in specific_overrides:
                            val = specific_overrides[(t_idx, r_idx)]
                            if isinstance(val, list):
                                answer = val[ff_i] if ff_i < len(val) else val[-1]
                            else:
                                answer = val
                                
                        # Generic fallbacks if needed
                        if not answer:
                            p_low = prompt_snippet.lower()
                            if "state/territory" in p_low:
                                answer = CANDIDATE_DETAILS["state"]
                            elif "organisation" in p_low or "workplace" in p_low:
                                answer = CANDIDATE_DETAILS["workplace"]
                            elif "supervisor" in p_low:
                                answer = "Sarah Jenkins, RN"
                            elif "date" in p_low:
                                answer = CANDIDATE_DETAILS["date"]
                            elif "signature" in p_low:
                                answer = CANDIDATE_DETAILS["candidate_name"]
                                
                    field_id = f"{unit_code}_TXT_{t_idx}_{r_idx}_{c_idx}_{ff_i+1}"
                    data["fields"].append({
                        "field_id": field_id,
                        "type": "text",
                        "section": section,
                        "question_ref": q_ref,
                        "prompt": prompt_snippet[:150],
                        "table_idx": t_idx,
                        "row_idx": r_idx,
                        "col_idx": c_idx,
                        "index_in_cell": ff_i,
                        "para_id": para_id,
                        "value": answer
                    })
                    
    # Write to YAML
    with open(output_yaml, 'w', encoding='utf-8') as f:
        yaml.dump(data, f, sort_keys=False, allow_unicode=True, width=1000)
        
    print(f"Generated {output_yaml}: {len(data['fields'])} fields extracted.")
    return data

def main():
    group3_configs = [
        ("CHCLEG001", "Assignment Materials-20260915 (7)/CHCLEG001-AWB-F-v1.1.docx", "answers_CHCLEG001.yaml"),
        ("HLTWHS002", "Assignment Materials-20260915 (9)/HLTWHS002-AWB-F-v1.0.docx", "answers_HLTWHS002.yaml"),
        ("CHCAGE011", "Assignment Materials-20260915 (10)/CHCAGE011-AWB-F-v1.0.docx", "answers_CHCAGE011.yaml"),
        ("CHCDIS011", "Assignment Materials-20260915 (12)/CHCDIS011-AWB-F-v1.1.docx", "answers_CHCDIS011.yaml"),
        ("CHCDIS012", "Assignment Materials-20260915 (13)/CHCDIS012-AWB-F-v1.0.docx", "answers_CHCDIS012.yaml"),
    ]
    
    for code, in_docx, out_yaml in group3_configs:
        process_file(code, in_docx, out_yaml)

if __name__ == "__main__":
    main()
