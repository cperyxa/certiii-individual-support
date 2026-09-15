import docx
from extractor_utils import extract_answer_items

def test_nested_match(awb_p, ag_p, t_idx):
    awb = docx.Document(awb_p)
    ag = docx.Document(ag_p)
    c = awb.tables[t_idx].rows[0].cells[0]
    for nt_i, nt in enumerate(c.tables):
        print(f"=== Table {t_idx} Nested Table {nt_i} ===")
        for r_i, r in enumerate(nt.rows[1:], 1):
            row_lbl = r.cells[0].text.strip().replace('\u2002', '')
            for c_i, cell in enumerate(r.cells[1:], 1):
                col_lbl = nt.rows[0].cells[c_i].text.strip().replace('\u2002', '')
                ffs = len(cell._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ffData'))
                if ffs > 0:
                    print(f"  Row: '{row_lbl}' | Col: '{col_lbl}' | ff={ffs}")
                    # Search AG
                    for ag_t in ag.tables:
                        for ag_r in ag_t.rows:
                            r_txt = " ".join([cell.text for cell in ag_r.cells])
                            if row_lbl.lower() in r_txt.lower() and col_lbl.lower()[:10] in r_txt.lower():
                                print(f"    -> MATCH IN AG: {[c.text.strip().replace(chr(10), ' ')[:50] for c in ag_r.cells]}")
                                break

print("Testing CHCCCS040 Table 13:")
test_nested_match("Assignment Materials-20260915 (3)/CHCCCS040-AWB-F-v1.0.docx", "Assignment Materials-20260915 (3)/CHCCCS040-AG-F-v1.0.docx", 13)

print("\nTesting HLTINF006 Table 13:")
test_nested_match("Assignment Materials-20260915 (8)/HLTINF006-AWB-F-v1.0.docx", "Assignment Materials-20260915 (8)/HLTINF006-AG-F-v1.0.docx", 13)

print("\nTesting CHCAGE013 Table 5:")
test_nested_match("Assignment Materials-20260915 (11)/CHCAGE013-AWB-F-v1.0.docx", "Assignment Materials-20260915 (11)/CHCAGE013-AG-F-v1.1.docx", 5)
