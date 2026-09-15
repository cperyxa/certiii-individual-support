import docx
import sys

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

units = [
    ('CHCCCS040', 'Assignment Materials-20260915 (3)/CHCCCS040-AG-F-v1.0.docx'),
    ('CHCCOM005', 'Assignment Materials-20260915 (5)/CHCCOM005-AG-F-v1.1.docx'),
    ('HLTINF006', 'Assignment Materials-20260915 (8)/HLTINF006-AG-F-v1.0.docx'),
    ('CHCAGE013', 'Assignment Materials-20260915 (11)/CHCAGE013-AG-F-v1.1.docx')
]

for code, ag_p in units:
    print(f"\n{'='*30} {code} STATE/NSW QUESTIONS {'='*30}")
    ag = docx.Document(ag_p)
    # Check Preliminary Task at Table 3
    t3 = ag.tables[3]
    t3_txt = " ".join([c.text for r in t3.rows for c in r.cells])
    print(f"Table 3 (Preliminary Task): {t3_txt[:150]}...")
    
    # Search for New South Wales in AG
    for i, t in enumerate(ag.tables):
        txt = " ".join([c.text for r in t.rows for c in r.cells])
        if "new south wales" in txt.lower():
            header = t.rows[0].cells[0].text.strip().replace('\n', ' ')[:100]
            print(f"AG Table {i}: {header}")
            # Find the NSW specific row or text
            for r in t.rows:
                for c in r.cells:
                    if "new south wales" in c.text.lower():
                        lines = [l.strip() for l in c.text.split('\n') if l.strip()]
                        for l in lines[:5]:
                            print(f"    * {l[:100]}")
                        break
