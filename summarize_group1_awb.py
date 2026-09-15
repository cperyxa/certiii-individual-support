import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

units = [
    ('CHCCCS040', 'Assignment Materials-20260915 (3)/CHCCCS040-AWB-F-v1.0.docx'),
    ('CHCCOM005', 'Assignment Materials-20260915 (5)/CHCCOM005-AWB-F-v1.1.docx'),
    ('HLTINF006', 'Assignment Materials-20260915 (8)/HLTINF006-AWB-F-v1.0.docx'),
    ('CHCAGE013', 'Assignment Materials-20260915 (11)/CHCAGE013-AWB-F-v1.0.docx')
]

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
W14_NS = 'http://schemas.microsoft.com/office/word/2010/wordml'

for code, awb_p in units:
    print(f"\n{'='*30} {code} AWB SUMMARY {'='*30}")
    doc = docx.Document(awb_p)
    total_ff = 0
    total_cb = 0
    for i, t in enumerate(doc.tables):
        ffs = t._element.findall(f'.//{{{W_NS}}}ffData')
        cbs = t._element.findall(f'.//{{{W14_NS}}}checkbox')
        if ffs or cbs:
            total_ff += len(ffs)
            total_cb += len(cbs)
            snippet = " ".join([c.text.strip().replace('\u2002', '') for r in t.rows for c in r.cells if c.text.strip()])[:90]
            print(f"Table {i:3d}: {len(ffs):2d} text, {len(cbs):2d} cb | {snippet}")
    print(f"TOTAL: {total_ff} text fields, {total_cb} checkboxes across {len(doc.tables)} tables")
