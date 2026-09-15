import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_tbl(p, t_idx):
    doc = docx.Document(p)
    t = doc.tables[t_idx]
    txt = " ".join([c.text.strip().replace('\u2002', '') for r in t.rows for c in r.cells if c.text.strip()])
    print(f"Table {t_idx} ({len(t.rows)} rows): {txt[:100]}")
    c = t.rows[0].cells[0]
    if c.tables:
        print(f"  Nested tables in cell 0: {len(c.tables)}")
        for nt in c.tables:
            for r in nt.rows[:3]:
                print(f"    {[cell.text.strip().replace(chr(8194), '')[:30] for cell in r.cells]}")

print("=== CHCAGE013 ===")
inspect_tbl("11. CHCAGE013 - Work effectively in aged care/CHCAGE013-AWB-F-v1.0.docx", 17)
inspect_tbl("11. CHCAGE013 - Work effectively in aged care/CHCAGE013-AWB-F-v1.0.docx", 48)

print("\n=== HLTINF006 ===")
inspect_tbl("8. HLTINF006 - Apply basic principles and practices of infection prevention and control/HLTINF006-AWB-F-v1.0.docx", 28)
inspect_tbl("8. HLTINF006 - Apply basic principles and practices of infection prevention and control/HLTINF006-AWB-F-v1.0.docx", 30)
inspect_tbl("8. HLTINF006 - Apply basic principles and practices of infection prevention and control/HLTINF006-AWB-F-v1.0.docx", 36)
